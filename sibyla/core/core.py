import sys
import os
import threading

import pdfplumber
import fitz  # PyMuPDF
from docx import Document
from docx.shared import Pt

from .translation import traduzir_texto, traduzir_blocos_pagina
from . import translation as _translation_module
from .image_utils import pagina_para_imagem_bytes
from .pdf_reader import extrair_blocos_pagina, extrair_imagens_da_pagina
from .word_writer import (
    inserir_imagem_no_doc,
    adicionar_pagina_no_doc,
    _encontrar_marcadores,
    _remover_pagina_do_doc,
    _mover_para_antes_de,
)
from .text_writer import salvar_txt, salvar_md
from .pdf_writer import salvar_pdf

_FORMATOS_TEXTO = {"txt", "md", "pdf"}


def _extrair_traduzir_pagina(num: int, pdf, fitz_doc,
                             lang_src: str, lang_dst: str,
                             modo_traducao: str = "bloco",
                             engine: str = "google",
                             api_key: str | None = None) -> dict:
    page   = pdf.pages[num - 1]
    blocos = extrair_blocos_pagina(page)
    imagens = extrair_imagens_da_pagina(fitz_doc, num - 1)

    if not blocos and not imagens:
        print("no text/images -> rendering page as image...", end=" ", flush=True)
        img_bytes = pagina_para_imagem_bytes(fitz_doc, num - 1)
        return {"num": num, "blocos": [], "imagens": [{"bytes": img_bytes, "y_top": 0}]}

    print(f"{len(blocos)} blocks. Translating...", end=" ", flush=True)
    if modo_traducao == "pagina":
        blocos_traduzidos = traduzir_blocos_pagina(blocos, lang_src, lang_dst, engine, api_key)
    else:
        blocos_traduzidos = [
            {**b, "texto": traduzir_texto(b["texto"], lang_src, lang_dst, engine, api_key)}
            for b in blocos
        ]
    print("OK")
    return {"num": num, "blocos": blocos_traduzidos, "imagens": imagens}


def _processar_pagina(num: int, pdf, fitz_doc, doc, modo: str,
                      ref_element=None,
                      lang_src: str = "en", lang_dst: str = "pt",
                      alinhamento: str = "original",
                      tamanho_fonte: int | None = None,
                      fonte: str | None = None,
                      modo_traducao: str = "bloco",
                      engine: str = "google",
                      api_key: str | None = None) -> None:
    page    = pdf.pages[num - 1]
    blocos  = extrair_blocos_pagina(page)
    imagens = extrair_imagens_da_pagina(fitz_doc, num - 1)
    count_before = len(doc.paragraphs)

    if not blocos and not imagens:
        print("no text/images -> rendering page as image...", end=" ", flush=True)
        img_bytes = pagina_para_imagem_bytes(fitz_doc, num - 1)
        p = doc.add_paragraph()
        run = p.add_run(f"— Página {num} —")
        run.font.hidden = True
        run.font.size = Pt(1)
        inserir_imagem_no_doc(doc, img_bytes)
        doc.add_page_break()
    else:
        print(f"{len(blocos)} blocks. Translating...", end=" ", flush=True)
        if modo_traducao == "pagina":
            blocos_traduzidos = traduzir_blocos_pagina(blocos, lang_src, lang_dst, engine, api_key)
        else:
            blocos_traduzidos = [
                {**bloco, "texto": traduzir_texto(bloco["texto"], lang_src, lang_dst, engine, api_key)}
                for bloco in blocos
            ]
        adicionar_pagina_no_doc(doc, num, blocos_traduzidos, imagens,
                                alinhamento=alinhamento,
                                tamanho_fonte=tamanho_fonte,
                                fonte=fonte)

    if modo == "replace" and ref_element is not None:
        _mover_para_antes_de(doc, count_before, ref_element)

    print("OK")


def processar_epub(epub_path: str, pag_ini: int, pag_fim: int, saida: str,
                   cancel_event: threading.Event | None,
                   lang_src: str, lang_dst: str,
                   fmt: str,
                   glossario: list[str] | None,
                   alinhamento: str,
                   tamanho_fonte: int | None,
                   fonte: str | None,
                   modo_traducao: str,
                   engine: str = "google",
                   api_key: str | None = None) -> None:
    import ebooklib
    from ebooklib import epub as _epub
    from bs4 import BeautifulSoup, NavigableString

    if glossario:
        _translation_module.NOMES_PROTEGIDOS.update(glossario)

    book = _epub.read_epub(epub_path, options={"ignore_ncx": True})

    spine_ids = {item_id for item_id, _ in book.spine}
    docs = [item for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT)
            if item.id in spine_ids]
    if not docs:
        docs = list(book.get_items_of_type(ebooklib.ITEM_DOCUMENT))

    total = len(docs)
    pag_fim = min(pag_fim, total)
    print(f"Total chapters: {total}\n")

    TRANSLATABLE = {"p", "h1", "h2", "h3", "h4", "h5", "h6",
                    "li", "td", "th", "caption", "blockquote", "dt", "dd"}

    def _translate_nodes(tag):
        for child in list(tag.children):
            if isinstance(child, NavigableString):
                if child.strip():
                    translated = traduzir_texto(str(child), lang_src, lang_dst, engine, api_key)
                    child.replace_with(translated or str(child))
            elif hasattr(child, "name") and child.name not in ("script", "style"):
                _translate_nodes(child)

    # For fmt != epub we accumulate translated text as paginas_data
    paginas_data: list[dict] = []

    for i, item in enumerate(docs):
        num = i + 1
        if num < pag_ini:
            continue
        if num > pag_fim:
            break
        if cancel_event and cancel_event.is_set():
            print("\nTranslation cancelled.")
            break

        print(f"[{num}/{pag_fim}] Chapter {num}...", end=" ", flush=True)

        raw_content = item.get_content()
        if not raw_content:
            print("(no content, skipped)")
            continue

        soup = BeautifulSoup(raw_content, "html.parser")
        elements = [el for el in soup.find_all(TRANSLATABLE)
                    if not el.find_parent(TRANSLATABLE) and el.get_text(strip=True)]

        if not elements:
            print("(empty)")
            if fmt != "epub":
                paginas_data.append({"num": num, "blocos": [], "imagens": []})
            continue

        if modo_traducao == "pagina":
            blocos = [{"texto": el.get_text(strip=True), "_el": el} for el in elements]
            traduzidos = traduzir_blocos_pagina(blocos, lang_src, lang_dst, engine, api_key)
            for b in traduzidos:
                b["_el"].clear()
                b["_el"].append(b["texto"] or "")
        else:
            for el in elements:
                _translate_nodes(el)

        item.set_content(str(soup).encode("utf-8"))

        if fmt != "epub":
            # Formato esperado pelos writers: {texto, tamanho, negrito, italico, x0, y_top, y_bot}
            blocos_txt = []
            for j, el in enumerate(elements):
                is_heading = el.name in ("h1", "h2", "h3", "h4", "h5", "h6")
                blocos_txt.append({
                    "texto":   el.get_text(strip=True),
                    "tamanho": tamanho_fonte or 11,
                    "negrito": is_heading,
                    "italico": False,
                    "x0":      0,
                    "y_top":   j * 20,
                    "y_bot":   j * 20 + 18,
                })
            paginas_data.append({"num": num, "blocos": blocos_txt, "imagens": []})
            if fmt == "txt":
                salvar_txt(paginas_data, saida)
            elif fmt == "md":
                salvar_md(paginas_data, saida)
            elif fmt == "pdf":
                salvar_pdf(paginas_data, saida, alinhamento=alinhamento,
                           tamanho_fonte=tamanho_fonte, fonte=fonte)
            elif fmt == "docx":
                _save_epub_as_docx(paginas_data, saida, alinhamento, tamanho_fonte, fonte)

        print("OK")

    if fmt == "epub" and not (cancel_event and cancel_event.is_set()):
        _salvar_epub(book, saida)

    print(f"\nDone! File saved at: {saida}")


def _save_epub_as_docx(paginas_data: list[dict], saida: str,
                       alinhamento: str, tamanho_fonte: int | None, fonte: str | None) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    _ALIGN = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Pt(72)

    for page in paginas_data:
        for bloco in page["blocos"]:
            is_heading = bloco.get("negrito", False)
            p = doc.add_heading(bloco["texto"], level=2) if is_heading else doc.add_paragraph(bloco["texto"])
            if not is_heading:
                align = _ALIGN.get(alinhamento)
                if align:
                    p.alignment = align
                for run in p.runs:
                    if tamanho_fonte:
                        run.font.size = Pt(tamanho_fonte)
                    if fonte:
                        run.font.name = fonte
        doc.add_page_break()
    doc.save(saida)


def _salvar_epub(book, saida: str) -> None:
    from ebooklib import epub as _epub

    def _fix_toc_uids(items, _n=[0]):
        for it in items:
            if isinstance(it, (tuple, list)):
                _fix_toc_uids(it, _n)
            elif hasattr(it, "uid") and it.uid is None:
                it.uid = f"toc{_n[0]}"
                _n[0] += 1

    _fix_toc_uids(book.toc)
    for _item in book.get_items():
        if _item.content is None:
            _item.content = b""
    _epub.write_epub(saida, book)


def processar_pdf_para_epub(pdf_path: str, pag_ini: int, pag_fim: int, saida: str,
                             cancel_event,
                             lang_src: str, lang_dst: str,
                             glossario: list[str] | None,
                             tamanho_fonte: int | None,
                             modo_traducao: str,
                             engine: str = "google",
                             api_key: str | None = None) -> None:
    import ebooklib
    from ebooklib import epub as _epub

    if glossario:
        _translation_module.NOMES_PROTEGIDOS.update(glossario)

    book = _epub.EpubBook()
    book.set_identifier(f"sibyla-{os.path.basename(pdf_path)}")
    book.set_title(os.path.splitext(os.path.basename(pdf_path))[0])
    book.set_language(lang_dst)

    fitz_doc = fitz.open(pdf_path)
    chapters = []

    with pdfplumber.open(pdf_path) as pdf:
        total = len(pdf.pages)
        pag_fim = min(pag_fim, total)
        print(f"Total pages: {total}\n")

        for num in range(pag_ini, pag_fim + 1):
            if cancel_event and cancel_event.is_set():
                print("\nTranslation cancelled.")
                break

            print(f"[{num}/{pag_fim}] Page {num}...", end=" ", flush=True)
            data = _extrair_traduzir_pagina(num, pdf, fitz_doc, lang_src, lang_dst,
                                            modo_traducao, engine, api_key)

            html_parts = []
            for bloco in data["blocos"]:
                text = bloco["texto"].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                if bloco.get("negrito"):
                    html_parts.append(f"<h2>{text}</h2>")
                elif bloco.get("italico"):
                    html_parts.append(f"<p><em>{text}</em></p>")
                else:
                    html_parts.append(f"<p>{text}</p>")

            for j, img in enumerate(data["imagens"]):
                img_id = f"img_p{num}_{j}"
                img_item = _epub.EpubImage()
                img_item.id = img_id
                img_item.file_name = f"images/{img_id}.png"
                img_item.media_type = "image/png"
                img_item.content = img["bytes"]
                book.add_item(img_item)
                html_parts.append(f'<p><img src="../images/{img_id}.png" alt=""/></p>')

            chapter = _epub.EpubHtml(
                title=f"Page {num}",
                file_name=f"page_{num:04d}.xhtml",
                lang=lang_dst,
            )
            chapter.content = f'<html><body>{"".join(html_parts)}</body></html>'.encode("utf-8")
            book.add_item(chapter)
            chapters.append(chapter)
            print("OK")

    fitz_doc.close()

    book.toc = chapters
    book.add_item(_epub.EpubNcx())
    book.add_item(_epub.EpubNav())
    book.spine = ["nav"] + chapters

    _salvar_epub(book, saida)
    print(f"\nDone! File saved at: {saida}")


def processar(pdf_path: str, pag_ini: int, pag_fim: int, saida: str,
              modo: str = "novo", arquivo_base: str | None = None,
              cancel_event: threading.Event | None = None,
              lang_src: str = "en", lang_dst: str = "pt",
              fmt: str = "docx",
              glossario: list[str] | None = None,
              alinhamento: str = "original",
              tamanho_fonte: int | None = None,
              fonte: str | None = None,
              modo_traducao: str = "bloco",
              engine: str = "google",
              api_key: str | None = None) -> None:

    ext = os.path.splitext(pdf_path)[1].lower()
    if ext == ".epub":
        processar_epub(pdf_path, pag_ini, pag_fim, saida, cancel_event,
                       lang_src, lang_dst, fmt, glossario,
                       alinhamento, tamanho_fonte, fonte, modo_traducao,
                       engine, api_key)
        return

    if fmt == "epub":
        processar_pdf_para_epub(pdf_path, pag_ini, pag_fim, saida, cancel_event,
                                lang_src, lang_dst, glossario,
                                tamanho_fonte, modo_traducao, engine, api_key)
        return

    print(f"\nFile    : {pdf_path}")
    print(f"Pages   : {pag_ini} to {pag_fim}")
    print(f"Mode    : {modo.upper()}")
    print(f"Lang    : {lang_src} -> {lang_dst}")
    print(f"Format  : {fmt.upper()}")
    print(f"Engine  : {engine.upper()}")
    print(f"Output  : {saida}\n")

    if glossario:
        _translation_module.NOMES_PROTEGIDOS.update(glossario)
        print(f"Glossary: {len(glossario)} protected terms\n")

    if fmt in _FORMATOS_TEXTO:
        fitz_doc = fitz.open(pdf_path)
        paginas_data: list[dict] = []
        with pdfplumber.open(pdf_path) as pdf:
            total   = len(pdf.pages)
            pag_fim = min(pag_fim, total)
            print(f"Total pages: {total}\n")
            for num in range(pag_ini, pag_fim + 1):
                if cancel_event is not None and cancel_event.is_set():
                    print("\nTranslation cancelled.")
                    break
                print(f"[{num}/{pag_fim}] Page {num}...", end=" ", flush=True)
                paginas_data.append(
                    _extrair_traduzir_pagina(num, pdf, fitz_doc, lang_src, lang_dst,
                                             modo_traducao, engine, api_key)
                )
                # Salva após cada página — garante arquivo parcial mesmo se cancelado
                if fmt == "txt":
                    salvar_txt(paginas_data, saida)
                elif fmt == "md":
                    salvar_md(paginas_data, saida)
                else:
                    salvar_pdf(paginas_data, saida, alinhamento=alinhamento,
                               tamanho_fonte=tamanho_fonte, fonte=fonte)
        fitz_doc.close()
        print(f"\nDone! File saved at: {saida}")
        return

    if modo in ("append", "replace"):
        if not arquivo_base or not os.path.isfile(arquivo_base):
            print(f"Error: base file '{arquivo_base}' not found.")
            sys.exit(1)
        doc = Document(arquivo_base)
        if modo == "replace":
            ranges = _encontrar_marcadores(doc)
            paginas_existentes = sorted(ranges.keys())
            print(
                f"Pages found in doc: "
                f"{paginas_existentes[0]}–{paginas_existentes[-1]}"
                if paginas_existentes else "  (no page markers found)"
            )
    else:
        doc = Document()
        for section in doc.sections:
            section.top_margin    = Pt(72)
            section.bottom_margin = Pt(72)
            section.left_margin   = Pt(72)
            section.right_margin  = Pt(72)

    fitz_doc = fitz.open(pdf_path)

    with pdfplumber.open(pdf_path) as pdf:
        total   = len(pdf.pages)
        pag_fim = min(pag_fim, total)
        print(f"Total pages: {total}\n")

        for num in range(pag_ini, pag_fim + 1):
            if cancel_event is not None and cancel_event.is_set():
                print("\nTranslation cancelled.")
                break
            print(f"[{num}/{pag_fim}] Page {num}...", end=" ", flush=True)

            if modo == "replace":
                ref_element = _remover_pagina_do_doc(doc, num)
                _processar_pagina(num, pdf, fitz_doc, doc, modo, ref_element,
                                  lang_src, lang_dst, alinhamento,
                                  tamanho_fonte, fonte, modo_traducao, engine, api_key)
            else:
                _processar_pagina(num, pdf, fitz_doc, doc, modo,
                                  lang_src=lang_src, lang_dst=lang_dst,
                                  alinhamento=alinhamento,
                                  tamanho_fonte=tamanho_fonte,
                                  fonte=fonte,
                                  modo_traducao=modo_traducao,
                                  engine=engine,
                                  api_key=api_key)

            # Salva após cada página — garante arquivo parcial mesmo se cancelado
            doc.save(saida)

    fitz_doc.close()
    print(f"\nDone! File saved at: {saida}")
